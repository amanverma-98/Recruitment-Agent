from io import BytesIO
from docx import Document


def generate_docx(questions,

    include_answers=False,

    include_explanations=False):

    document = Document()

    document.add_heading("Question Bank",level=1)
    for index, question in enumerate(questions, start=1):
        document.add_heading(f"{index}. {question.question_text}", level=2)
        document.add_paragraph(f"Topic: {question.topic}")
        document.add_paragraph(f"Difficulty: {question.difficulty}")
        document.add_paragraph("Options:")
        for option in question.options:
            document.add_paragraph(option, style="List Bullet")
        if include_answers:
            document.add_paragraph(f"Correct Answer: {question.correct_option}")
        if include_explanations:
            document.add_paragraph(question.explanation)
        document.add_page_break()

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)

    return buffer